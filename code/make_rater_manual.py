# make_rater_manual.py — 評価者マニュアル（Rater Manual）を ../rater_package/ に生成する
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "rater_package", "評価者マニュアル_Rater_Manual.docx")
JP = "游ゴシック"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.2))
for name in ("Normal", "List Bullet", "List Number", "Heading 1", "Heading 2", "Title"):
    s = doc.styles[name]
    s.font.name = "Times New Roman"
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), JP)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Title", 17)):
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles[name].font.bold = True


def p(text="", bold=None, size=None, align=None):
    para = doc.add_paragraph()
    if bold:
        r = para.add_run(bold); r.bold = True
    r = para.add_run(text)
    if size:
        r.font.size = Pt(size)
    if align:
        para.alignment = align
    return para


def bullet(text, bold=None, style="List Bullet"):
    para = doc.add_paragraph(style=style)
    if bold:
        r = para.add_run(bold); r.bold = True
    para.add_run(text)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def table(rows, widths, header_fill="D9D9D9"):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.text = ""
            r = c.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            if i == 0:
                r.bold = True
                shade(c, header_fill)
    doc.add_paragraph()
    return t


doc.add_paragraph("AI生成研究原稿の盲検評価", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
p("評価者マニュアル（Rater Manual）", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
p("依頼者：大槻 明（日本大学 経済学部）　連絡先：otsuki.akira@nihon-u.ac.jp", size=9.5,
  align=WD_ALIGN_PARAGRAPH.CENTER)

# ---------------------------------------------------------------- 0
doc.add_heading("このマニュアルの要点（最初にお読みください）", level=1)
table([
    ["項目", "内容"],
    ["お願いしたいこと", "AIが自動生成した英語の研究原稿 20本を読み、3つの観点（技術的妥当性・再現性・重要性）を1〜5点で採点し、評価シートに記入して返送していただくこと"],
    ["所要時間の目安", "1本あたり約12分（読む10分＋採点2分）、合計で約4時間。何日かに分けて構いません"],
    ["提出期限", "2026年10月　　日（　）　※依頼者が記入"],
    ["提出物", "記入済みの評価シート（Excelファイル）1つだけ"],
    ["最も大切な約束", "①他の評価者と相談しない　②原稿の作られ方を推測しようとしない　③ChatGPTなどの生成AIに原稿を読ませたり採点させたりしない　④原稿を第三者に渡さない"],
], [3.5, 12.5])

# ---------------------------------------------------------------- 1
doc.add_heading("1. この評価の目的", level=1)
p("研究テーマ（仮説）を自動で作り、そこから研究原稿を自動で書くAIシステムを研究しています。原稿は、いくつかの異なる条件で生成しました。"
  "どの条件で生成したものが、専門家の目から見て質が高いのかを調べるのが、この評価の目的です。")
p("どの原稿がどの条件で作られたかは、評価が終わるまで評価者の方にはお知らせしません（盲検評価）。条件の内容も、評価の公平性を保つため、"
  "評価終了後にご説明します。原稿はすべて同じAIが、同じ指示文で書いています。")
p("評価結果は、学術誌（MDPI Informatics）に投稿中の論文の改訂版に、統計的に集計した形で掲載します。"
  "評価者個人の点数やコメントが、氏名と結び付く形で公開されることはありません。")

# ---------------------------------------------------------------- 2
doc.add_heading("2. お送りするファイル", level=1)
table([
    ["ファイル", "内容"],
    ["評価者マニュアル_Rater_Manual.docx", "このマニュアル"],
    ["manuscripts\\P00_practice.docx", "練習用の原稿（1本）。分析には使いません"],
    ["manuscripts\\M0＊＊.docx（20ファイル）", "評価していただく原稿（20本）。番号はランダムで飛び飛びです。番号や順番に意味はありません"],
    ["評価シート_R＊.xlsx（R1またはR2）", "採点を記入するシート。評価者ごとに原稿を読む順番が違います"],
], [6.5, 9.5])

# ---------------------------------------------------------------- 3
doc.add_heading("3. 評価の手順", level=1)
for t in [
    "このマニュアルの「4. 採点基準」を一度通して読みます。",
    "練習用の原稿 P00_practice.docx を読み、評価シートの「練習」シートに採点してみます。採点基準の感覚をつかむためのもので、分析には使いません。",
    "評価シートの「評価」シートに書かれた順番（position 1, 2, 3, …）で原稿を読みます。読む順番は評価者ごとに変えてあります。順番どおりに読んでください。",
    "1本読み終えるごとに、その原稿の行に、3観点の点数（1〜5）を記入します。コメントは任意です。",
    "一度つけた点数は、後で他の原稿と見比べて直しても構いません。ただし、全部読み終えてから一斉に付け直すことはせず、原則として読んだ直後の判断を記入してください。",
    "20本すべて記入したら、ファイル名を変えずに、依頼者へメールで返送してください。",
]:
    bullet(t, style="List Number")

# ---------------------------------------------------------------- 4
doc.add_heading("4. 採点基準（1〜5点）", level=1)
p("次の3つの観点を、それぞれ1〜5の整数で採点してください。3は「平均的な研究提案として普通」の水準です。"
  "迷ったときは、下の説明に一番近い点数を選んでください。")
crit = [
    ("4.1 技術的妥当性（Technical soundness）",
     "提案手法、数式、論理の流れが正しく一貫しているか。手法と問題が本当に結び付いているか。誤りや論理の飛躍がないか。",
     ["1：重大な誤りがある、または手法が問題の解決につながっていない",
      "2：誤りや飛躍が目立ち、主張の根拠が弱い",
      "3：おおむね妥当だが、あいまいな点や小さな誤りがある",
      "4：大きな問題はなく、論理が通っている",
      "5：一貫していて正確。専門家として納得できる"]),
    ("4.2 再現性（Reproducibility）",
     "本文の記述だけで、提案手法を実装し、実験計画を実行できるか（手順・設定・評価方法が具体的に書かれているか）。",
     ["1：何をすればよいか分からず、実装できない",
      "2：方針は分かるが、重要な部分がほとんど書かれていない",
      "3：主要部分は実装できるが、重要な詳細がいくつか欠けている",
      "4：少し補えば実装・実行できる",
      "5：十分に具体的で、そのまま実装・実行できる"]),
    ("4.3 重要性（Significance）",
     "研究として追究する価値があるか。成功した場合の影響の大きさや、応用できる範囲の広さ。",
     ["1：追究する価値がほとんどない",
      "2：価値はあるが、非常に限定的",
      "3：一定の価値があり、関連分野の研究者には有用",
      "4：多くの研究者や応用分野にとって価値がある",
      "5：分野全体に大きな影響を与えうる"]),
]
for title, q, anchors in crit:
    doc.add_heading(title, level=2)
    p(q, bold="問い：")
    for a in anchors:
        bullet(a)

doc.add_heading("4.4 採点するときの注意", level=2)
for b, t in [
    ("アイデアが新しいかどうかは採点の対象外です。", "研究分野の動向に詳しくなくても評価できるよう、新しさ（独創性）は今回の評価項目に含めていません。上の3つの観点だけで採点してください。"),
    ("実験結果がないことでは減点しないでください。", "どの原稿も、実験は「今後の計画」として書くよう指示されています。計画として妥当かどうかで評価してください。"),
    ("引用文献がないことでは減点しないでください。", "どの原稿も、架空の文献を作らないよう、具体的な引用を避けて書くよう指示されています。"),
    ("文章のうまさ、長さ、見た目は評価に含めないでください。", "英語の自然さや体裁ではなく、内容で評価してください。"),
    ("数値が出てきた場合", "原稿内に具体的な実験結果の数値が書かれていても、それは実際の実験に基づくものではありません。その数値を根拠に高く評価しないでください。また、根拠のない数値を結果として書いている場合は、技術的妥当性の評価に反映して構いません。"),
    ("原稿どうしは独立に評価してください。", "似たテーマの原稿があっても、それぞれ単独で採点してください。"),
]:
    bullet(t, bold=b + " ")

doc.add_heading("4.5 AIが専門でない評価者の方へ", level=2)
p("評価者の皆様の専門分野はAI・機械学習とは限りません。それで問題ありません。")
for b, t in [
    ("技術的妥当性・再現性は、ご自身の専門の感覚で評価してください。", "数式や論理の整合性、手法の記述の具体性は、分野を問わず判断できます。"),
    ("重要性は、研究として追究する価値があるかを、分かる範囲で判断してください。", "応用の広さ・社会的な意義など、ご自身の視点からの判断で構いません。"),
    ("専門用語が分からない場合は、", "文献検索や辞書で用語の意味を調べて構いません（生成AIは使わないでください）。"),
]:
    bullet(t, bold=b + " ")

doc.add_heading("4.6 コメント欄（任意）", level=2)
p("採点の理由や気づいた点を自由に記入してください。特に、1点や5点を付けた場合は一言書いていただけると助かります。")

# ---------------------------------------------------------------- 5
doc.add_heading("5. 守っていただきたいこと", level=1)
for b, t in [
    ("独立して評価してください。", "他の評価者とは、評価が終わるまで原稿や点数について話さないでください。"),
    ("原稿の作られ方を推測しないでください。", "どの条件で作られたかを推測して点数を調整することは避け、内容だけで評価してください。"),
    ("生成AIを使わないでください。", "ChatGPT、Gemini、Claude などに原稿を読ませたり、要約・採点させたりしないでください。"),
    ("原稿を外部に出さないでください。", "原稿は未発表の研究資料です。第三者に渡したり、インターネット上にアップロードしたりしないでください。評価終了後は削除していただいて構いません。"),
    ("利益相反がないことを確認してください。", "この研究に、著者・共同研究者・助言者として関わっていないことが評価者の条件です。該当する可能性がある場合は、評価を始める前に依頼者へご連絡ください。"),
]:
    bullet(t, bold=b + " ")

# ---------------------------------------------------------------- 6
doc.add_heading("6. よくある質問", level=1)
for q, a in [
    ("専門外のテーマの原稿があります。", "そのまま、分かる範囲で採点してください。"),
    ("数式がおかしい、または意味が通らない箇所があります。", "技術的妥当性の点数に反映してください。必要ならコメント欄に一言書いてください。"),
    ("原稿の途中で文章が切れています。", "読める範囲で採点し、コメント欄に「途中で切れている」と書いてください。"),
    ("同じようなテーマの原稿が複数あります。", "生成の性質上ありえます。それぞれ独立に採点してください。"),
    ("全部を一度に読む時間がありません。", "何日かに分けて構いません。その場合も、シートに書かれた順番を守ってください。"),
]:
    p(a, bold=f"Q. {q}\nA. ")

# ---------------------------------------------------------------- 7
doc.add_heading("7. 結果の扱いと謝辞", level=1)
p("採点結果は、条件ごとの平均値や評価者間の一致度として統計的に処理し、論文に掲載します。"
  "論文の Acknowledgments（謝辞）にお名前を記載してよいかどうかは、評価シートの「評価者情報」シートでお知らせください。"
  "掲載を希望されない場合は、お名前は記載しません。")
p("ご協力に心より感謝いたします。ご不明な点があれば、いつでも依頼者までご連絡ください。")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved", OUT)

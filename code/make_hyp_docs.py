# make_hyp_docs.py — 仮説の盲検評価（実験5）用の文書を生成する
#   ../rater_package_hypotheses/評価者マニュアル_仮説の評価.docx   （評価者用：これ1つで完結）
#   ../盲検評価_実施マニュアル（大槻先生用）.docx                  （依頼者用：評価者には送らない）
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.join(HERE, "..")
JP = "游ゴシック"


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    for name in ("Normal", "List Bullet", "List Number", "Heading 1", "Heading 2", "Title"):
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:eastAsia"), JP)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Title", 18)):
        st = doc.styles[name]
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0, 0, 0)
    return doc


class W:
    def __init__(self, doc):
        self.d = doc

    def p(self, text="", bold=None, size=None, align=None):
        para = self.d.add_paragraph()
        if bold:
            para.add_run(bold).bold = True
        r = para.add_run(text)
        if size:
            r.font.size = Pt(size)
        if align:
            para.alignment = align
        return para

    def bullet(self, text, bold=None, style="List Bullet"):
        para = self.d.add_paragraph(style=style)
        if bold:
            para.add_run(bold).bold = True
        para.add_run(text)

    @staticmethod
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
        tcPr.append(sh)

    def table(self, rows, widths, fills=None):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
                r = c.paragraphs[0].add_run(val); r.font.size = Pt(9.5)
                if i == 0:
                    r.bold = True; self.shade(c, "D9D9D9")
                elif fills and fills.get((i, j)):
                    self.shade(c, fills[(i, j)])
        self.d.add_paragraph()

    def box(self, lines, fill):
        t = self.d.add_table(rows=1, cols=1); t.style = "Table Grid"
        c = t.cell(0, 0); c.width = Cm(17); self.shade(c, fill); c.text = ""
        for k, (b, txt) in enumerate(lines):
            para = c.paragraphs[0] if k == 0 else c.add_paragraph()
            if b:
                r = para.add_run(b); r.bold = True; r.font.size = Pt(10)
            r = para.add_run(txt); r.font.size = Pt(10)
        self.d.add_paragraph()

    def mail(self, title, body):
        self.d.add_heading(title, level=2)
        t = self.d.add_table(rows=1, cols=1); t.style = "Table Grid"
        c = t.cell(0, 0); c.width = Cm(17); self.shade(c, "F2F2F2"); c.text = ""
        for k, line in enumerate(body.split("\n")):
            para = c.paragraphs[0] if k == 0 else c.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            r = para.add_run(line); r.font.size = Pt(9.5)
        self.d.add_paragraph()

    def check(self, items):
        for t in items:
            para = self.d.add_paragraph(); para.add_run("☐  ").font.size = Pt(12); para.add_run(t)


USEFUL = [
    ("① 課題の重要性：", "その仮説が取り組む課題（問題）は、解決する価値のある重要なものか。"),
    ("② 方法の妥当性：", "提案された方法は、その課題の解決や目標の達成に役立つ見込みがあるか。方法と課題・目標が論理的につながっているか。"),
    ("③ 実行可能性：", "実際に研究として取り組めそうか（現実的な手段で検証できそうか）。"),
]


# ====================================================================== rater manual
def rater_manual():
    doc = new_doc(); w = W(doc)
    doc.add_paragraph("研究仮説の盲検評価　評価者マニュアル", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    w.p("依頼者：大槻 明（日本大学 経済学部）　連絡先：otsuki.akira@nihon-u.ac.jp", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_heading("このマニュアルの要点（最初にお読みください）", level=1)
    w.table([
        ["項目", "内容"],
        ["お願いしたいこと", "AIの研究についての「研究仮説」が3つずつ並んだ組（A・B・C）を30組読み、各組で「有用性」と「新しさ」のそれぞれについて、最も高い仮説と最も低い仮説を1つずつ選んで、評価シートに記入していただくこと"],
        ["所要時間の目安", "1組あたり約2〜3分（3つの仮説はそれぞれ数行です）、合計で約1時間15分。何回かに分けて構いません"],
        ["提出期限", "2026年10月　　日（　）"],
        ["提出物", "記入済みの評価シート（Excelファイル）1つだけ"],
        ["大切な約束", "①他の評価者と相談しない　②仮説の作られ方を推測しようとしない　③ChatGPTなどの生成AIを使わない　④資料を第三者に渡さない"],
    ], [3.5, 13.5])

    doc.add_heading("1. この評価の目的", level=1)
    w.p("研究テーマ（研究仮説）を自動で作るAIシステムを研究しています。仮説はいくつかの異なる方法で作りました。"
        "どの方法で作った仮説が、研究者の目から見て有用で新しいかを調べるのが、この評価の目的です。")
    w.p("どの仮説がどの方法で作られたかは、評価が終わるまでお知らせしません（盲検評価）。"
        "どの仮説も、同じ書式・同じくらいの長さにそろえてあります。")
    w.p("評価結果は、学術誌（MDPI Informatics）に投稿中の論文の改訂版に、統計的に集計した形で掲載します。"
        "個人の回答が氏名と結び付く形で公開されることはありません。")

    doc.add_heading("2. お送りしたファイル", level=1)
    w.table([
        ["ファイル", "内容"],
        ["評価者マニュアル_仮説の評価.docx", "このマニュアル"],
        ["仮説冊子_R＊.docx", "評価していただく仮説の冊子。最初に練習用の1組（P0）、続いて本番の30組（No. 1〜30）が載っています"],
        ["評価シート_仮説_R＊.xlsx", "回答を記入するシート（「練習」「評価」「評価者情報」の3つのシート）"],
    ], [6.0, 11.0])

    doc.add_heading("3. 評価の手順", level=1)
    for t in ["このマニュアルの「4. 評価の基準」を読みます。",
              "冊子の練習用の組（P0）を読み、評価シートの「練習」シートに回答してみます（分析には使いません）。",
              "冊子の No. 1 から順に、1組ずつ3つの仮説（A・B・C）を読みます。",
              "各組について、評価シートの「評価」シートの同じ番号の行に、「最も有用」「最も有用でない」「最も新しい」「最も新しくない」を A/B/C からプルダウンで選びます。同じ観点で同じ記号を両方に選ぶと赤く表示されます。コメントは任意です。",
              "30組すべて記入したら、「評価者情報」シートを記入し、ファイル名を変えずに依頼者へメールで返送してください。"]:
        w.bullet(t, style="List Number")

    doc.add_heading("4. 評価の基準：有用性と新しさ", level=1)
    w.p("各組の3つの仮説を比べ、次の2つの観点それぞれについて、最も高いもの・最も低いものを1つずつ選んでください。"
        "2つの観点は別々に判断してください（有用だから新しい、新しいから有用、とは考えないでください）。")
    doc.add_heading("観点1　有用性", level=2)
    w.p("研究テーマとして取り組む価値があるか。次の3点を総合して判断します。")
    for b, t in USEFUL:
        w.bullet(t, bold=b)
    doc.add_heading("観点2　新しさ", level=2)
    w.p("その仮説のアイデア（どの方法を、どの課題に、何のために使うか）が、あなたが知っている研究や一般的な考え方と比べて、"
        "どれだけ新しく、意外性があるか。")
    w.bullet("AI分野の最新の研究動向をすべて知っている必要はありません。あなたの知識の範囲で、「よく聞く組み合わせか」「思いつきにくい組み合わせか」を判断してください。")
    w.bullet("実現できるかどうか（有用性の③）は、新しさの判断には含めないでください。")
    w.box([("どちらの観点でも判断に含めないこと", ""),
           ("", "・文章の上手さ・長さ・専門用語の多さ：難しい言葉が多いから有用・新しい、とは判断しないでください。"),
           ("", "・A・B・C の位置や組の番号：意味はありません（ランダムに並べています）。")], "FFF2CC")
    w.box([("評価者の専門性について", ""),
           ("", "AIの細かな研究動向を知っている必要はありません。研究者としての視点から判断していただければ十分です。"
                "新しさの判断が難しい組があっても、あなたの知識の範囲で選んでください。")], "DDEBF7")
    w.p("3つが同じくらいに思える組でも、わずかな差で構いませんので、必ず1つずつ選んでください。", size=9.5)

    doc.add_heading("5. 守っていただきたいこと", level=1)
    for b, t in [("独立して評価する：", "他の評価者とは、評価が終わるまで内容や回答について話さないでください。"),
                 ("内容だけで評価する：", "仮説がどの方法で作られたかを推測して回答を調整しないでください。"),
                 ("生成AIを使わない：", "ChatGPT・Gemini・Claude などに仮説を読ませたり、判断させたりしないでください。"),
                 ("資料を外部に出さない：", "未発表の研究資料です。第三者に渡したり、インターネットにアップロードしたりしないでください。評価終了後は削除して構いません。"),
                 ("利益相反がないことを確認する：", "この研究に、著者・共同研究者・助言者として関わっていないことが条件です。")]:
        w.bullet(t, bold=b, style="List Number")

    doc.add_heading("6. よくある質問", level=1)
    for q, a in [("どの仮説がどの方法で作られたのか教えてもらえますか。", "公平な評価のため、評価がすべて終わってからご説明します。"),
                 ("3つとも同じくらい良い（悪い）と思います。", "わずかな差で構いませんので、必ず1つずつ選んでください。"),
                 ("新しさの判断に自信がありません。", "あなたの知識の範囲での判断で構いません。「よく聞く組み合わせか、思いつきにくい組み合わせか」を目安にしてください。"),
                 ("知らない専門用語があります。", "文献検索や辞書で調べて構いません（生成AIは使わないでください）。分からないままでも、分かる範囲で判断していただければ十分です。"),
                 ("似た内容の仮説が別の組にもあります。", "ありえます。組ごとに独立に判断してください。"),
                 ("プルダウンが動きません。", "Excel 上部の「編集を有効にする」を押してください。それでも動かない場合は、A・B・C を直接入力してください。")]:
        w.p(a, bold=f"Q. {q}\nA. ")

    doc.add_heading("7. 提出前のチェックリスト", level=1)
    w.check(["「評価」シートの30組すべてに4つの欄（最も有用・最も有用でない・最も新しい・最も新しくない）を記入した（赤く表示された欄がない）",
             "「評価者情報」シートを記入した", "ファイル名を変えていない",
             "評価が終わるまで、他の評価者と相談していない／生成AIを使っていない"])
    out = os.path.join(REV, "rater_package_hypotheses", "評価者マニュアル_仮説の評価.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out); print("saved", out)


# ====================================================================== author guide
def author_guide():
    doc = new_doc(); w = W(doc)
    ME, RATER, CLAUDE = "大槻先生", "評価者", "Claude"
    Y, B, G = "FFF2CC", "DDEBF7", "E2EFDA"
    doc.add_paragraph("盲検評価 実施マニュアル（大槻先生用）", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    w.p("— 研究仮説の盲検評価：大槻先生と評価者2名が、いつ・何をすればよいか —", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    w.box([("このマニュアルは評価者には送らないでください。", ""),
           ("", "対応表ファイルの場所や、仮説の作り方（3つの条件）など、評価者に知られてはいけない情報が含まれます。"
                "評価者には rater_package_hypotheses フォルダの中のもの（評価者マニュアル・仮説冊子・評価シート）だけを送ります。")], "FCE4D6")

    doc.add_heading("1. 評価の内容と全体の流れ", level=1)
    w.table([
        ["項目", "内容"],
        ["評価するもの", "研究仮説30組。各組は、GAが選んだ仮説・候補から無作為に選んだ仮説・LLM自身が提案した仮説の3案（A・B・C、どれがどれかは伏せる）。3案とも同じAIで同じ書式・同じ長さに書き直してある"],
        ["評価者がすること", "各組で、有用性と新しさのそれぞれについて「最も高い」「最も低い」を1つずつ選ぶ"],
        ["評価者の人数と負担", "2名。2名とも同じ30組を評価する（読む順番は評価者ごとに違う）。1人あたり約1時間15分"],
        ["分かること", "どの方法で作った仮説が有用と判断されるか（条件間の比較）、評価者2名の判断の一致度"],
        ["新しさ（独創性）の扱い", "評価者による新しさの順位づけ（評価者はAI関連分野の研究者だが、AI研究と比べた新しさの判断には専門性が十分ではないため、参考の位置づけ）と、デスクトップPCでの大きめのLLMによる同じ30組の順位づけで示す。なお、自動の新規性指標は、OpenReviewの査読データ（ICLR 2023〜2025）との相関がほぼゼロで、独創性の根拠には使えないことが分かった"],
    ], [4.0, 13.0])
    w.table([
        ["担当", "役割"],
        [ME, "評価者2名に依頼する、資料を送る、催促する、返送された評価シートを Claude に渡す。並行して、デスクトップPCでLLM採点を実行する（別紙の手順書）"],
        [RATER + "（2名）", "マニュアルを読み、仮説30組を評価して、評価シートを返送する（約1時間15分）"],
        [CLAUDE, "評価用資料の作成、OpenReviewでの指標の検証、返送された評価の集計・統計分析、論文と回答書への反映"],
    ], [3.5, 13.5], fills={(1, 0): Y, (2, 0): G, (3, 0): B})
    w.p("日程（目安）", bold="")
    sched = [("10/2（金）", CLAUDE, "評価用資料（rater_package_hypotheses）の完成"),
             ("10/2（金）〜3（土）", ME, "評価者2名に依頼し（文例A）、承諾をもらったら資料を送る（文例B）"),
             ("10/3（土）〜10/7（水）", RATER, "評価（約1時間15分）"),
             ("10/6（火）", ME, "進み具合の確認（文例C）"),
             ("10/8（木）", RATER, "評価シートを返送（提出期限。最終期限は10/9正午）"),
             ("10/2〜10/8の間", ME, "デスクトップPCでLLM採点を実行し、結果をノートPCに戻す（別紙：手順書_デスクトップでのLLM採点）"),
             ("10/8〜10", CLAUDE, "集計・分析、論文（実験4・5、独創性の検証、要旨・考察）と回答書の更新"),
             ("10/10〜11", ME, "最終確認して再投稿。評価者にお礼と種明かし（文例D）")]
    w.table([["日付", "誰が", "やること"]] + [list(s) for s in sched], [3.8, 2.4, 10.8],
            fills={(i, 1): {CLAUDE: B, ME: Y, RATER: G}[s[1]] for i, s in enumerate(sched, 1)})

    doc.add_heading("2. 大槻先生の作業", level=1)
    doc.add_heading("STEP 1　評価者2名に依頼する", level=2)
    for t in ["AIに関連する分野（経営工学・計量経済学・流体力学など）の研究者で、英語の文章を読める方",
              "この研究に、著者・共同研究者・助言者として関わっていない方",
              "約1時間15分の作業を、10/8までに確保できる方"]:
        w.bullet(t)
    w.p("論文には、評価者を「AIに関連する分野の研究者であり、仮説の有用性は判断できるが、既存のAI研究と比べた新しさの評価には"
        "専門性が十分ではないため、新しさの評価は参考として扱う」と記載します。評価者が確定したら、お二人の専門分野を Claude に伝えてください。", size=9.5)

    doc.add_heading("STEP 2　資料を送る", level=2)
    w.table([
        ["送るもの", "評価者R1", "評価者R2"],
        [r"rater_package_hypotheses\R1 フォルダ（評価者マニュアル・仮説冊子_R1・評価シート_仮説_R1）", "○", "×"],
        [r"rater_package_hypotheses\R2 フォルダ（評価者マニュアル・仮説冊子_R2・評価シート_仮説_R2）", "×", "○"],
    ], [11.0, 3.0, 3.0])
    for t in ["評価者マニュアル（R1・R2 両フォルダ内）の冒頭の表の「提出期限」欄に日付を書き込んで保存する。",
              "R1 フォルダを右クリック →「送る」→「圧縮(zip形式)フォルダー」でZIPにする。R2 も同様にする。",
              "大学のファイル共有サービス、またはメール添付で送る（文例B）。"]:
        w.bullet(t, style="List Number")
    w.box([("絶対に送らないもの　", r"results\hypothesis_eval\_TRIPLET_KEY_do_not_share.csv（どの仮説がどの条件かの対応表）。"
             "これが評価者に渡ると盲検ではなくなります。"),
           ("評価者に伝えないこと　", "3つの条件の内容（GA／無作為／LLM自身の提案）と、2名が同じ30組を評価していること。")], "FCE4D6")

    doc.add_heading("STEP 3　進み具合を確認する（10/6頃）", level=2)
    w.p("文例Cで確認します。評価基準について質問された場合は、評価者マニュアルの「4. 評価の基準」を案内し、具体的な答え"
        "（「この組ならAが良い」など）は示さないでください。必要なら、2名に同じ文面で回答します。")

    doc.add_heading("STEP 4　返送された評価シートを Claude に渡す（10/8〜9）", level=2)
    for t in [r"返送された 評価シート_仮説_R1.xlsx と 評価シート_仮説_R2.xlsx を、Revision1\returned_ratings_hypotheses\ フォルダ（なければ作成）に保存する。ファイル名は変えない。",
              "Claude に「returned_ratings_hypotheses に評価シートを入れたので集計して論文に反映して」と伝える。"]:
        w.bullet(t, style="List Number")

    doc.add_heading("STEP 5　お礼と種明かし（再投稿後）", level=2)
    w.p("文例Dでお礼と条件の説明をします。謝辞への氏名掲載に同意した方（評価シートの「評価者情報」で「可」）は、論文の"
        "Acknowledgments に記載します（Claude が反映します）。")

    doc.add_heading("3. 評価者の作業（要約）", level=1)
    w.table([["順番", "やること", "目安"],
             ["1", "評価者マニュアルを読む", "10分"],
             ["2", "練習用の組（P0）を読み、「練習」シートに回答する", "5分"],
             ["3", "冊子の No. 1〜30 を順に読み、各組で有用性と新しさの最も高い・低い仮説を選ぶ", "2〜3分 × 30組"],
             ["4", "「評価者情報」シートを記入し、評価シートを返送する", "5分"]], [1.5, 12.5, 3.0])

    doc.add_heading("4. メール文例", level=1)
    w.mail("文例A　依頼メール", "件名：研究評価へのご協力のお願い（研究仮説の盲検評価）\n\n〇〇先生\n\n日本大学経済学部の大槻です。\n"
           "現在、AIが研究テーマ（研究仮説）を自動で作るシステムについての論文を学術誌（MDPI Informatics）に投稿しており、"
           "査読者から研究者による評価を求められています。つきましては、研究仮説が3つずつ並んだ組を30組読み、各組で最も有用な仮説と"
           "最も有用でない仮説、最も新しい仮説と最も新しくない仮説を選んでいただけないでしょうか。\n\n・作業量：合計約1時間15分（各仮説は数行です。何回かに分けて構いません）\n"
           "・期間：10月3日頃に資料をお送りし、10月8日（木）までにご返送いただく予定です\n・条件：本研究に関わっていない方にお願いしています\n"
           "・お礼：ご希望により、論文の謝辞にお名前を記載させていただきます\n\nご検討いただけますと幸いです。\n\n大槻 明")
    w.mail("文例B　資料送付メール", "件名：【資料送付】研究仮説の盲検評価\n\n〇〇先生\n\nこのたびはご協力をお引き受けいただき、ありがとうございます。"
           "評価用の資料をお送りします。\n（ZIPファイル添付、または共有リンク：＿＿＿＿）\n\n最初に「評価者マニュアル」をお読みください。"
           "練習用の組で感覚をつかんでいただいた後、冊子の No. 1 から順に30組を評価してください。\n\n提出期限：10月＿日（＿）\n"
           "提出物：記入済みの評価シート（Excelファイル）\n\n公平な評価のため、評価が終わるまで、他の評価者の方との相談や生成AIの利用は"
           "お控えください。仮説の作り方については、評価終了後にご説明いたします。\n\n大槻 明")
    w.mail("文例C　進み具合の確認メール", "件名：研究仮説の評価の進み具合について\n\n〇〇先生\n\n評価作業にご協力いただき、ありがとうございます。"
           "提出期限の10月＿日が近づいてまいりましたので、ご連絡いたしました。ご不明な点や、期限に間に合わない可能性がございましたら、"
           "お気軽にお知らせください。\n\n大槻 明")
    w.mail("文例D　お礼と種明かしのメール", "件名：研究仮説の評価のお礼\n\n〇〇先生\n\nこのたびは評価作業にご協力いただき、誠にありがとうございました。"
           "おかげさまで論文を再投稿することができました。\n評価いただいた各組の3つの仮説は、①遺伝的アルゴリズムが選んだ仮説、"
           "②同じ候補から無作為に選んだ仮説、③AI（大規模言語モデル）自身が提案した仮説を、同じAIで同じ書式に書き直したものでした。"
           "結果の概要は＿＿＿＿のとおりです。\n\n改めて御礼申し上げます。\n\n大槻 明")

    doc.add_heading("5. 困ったときは", level=1)
    w.table([["状況", "対応"],
             ["評価者のうち1名が辞退・期限に間に合わない", "1名分だけでも条件間の比較はできます（一致度は計算できなくなります）。Claude に伝えてください"],
             ["評価者から「どの条件か」を聞かれた", "「評価が終わってから説明します」と答える"],
             ["評価シートのプルダウンが動かない", "「編集を有効にする」を押してもらう。だめなら A・B・C を直接入力してもらう"]], [5.5, 11.5])

    doc.add_heading("6. チェックリスト", level=1)
    w.check(["評価者2名から承諾をもらった（研究に関わっていない方である）", "評価者マニュアルに提出期限を記入した",
             "R1 には R1 フォルダ、R2 には R2 フォルダのZIPを送った", "_TRIPLET_KEY_do_not_share.csv を送っていない",
             "10/6頃に進み具合を確認した", "返送された2ファイルを returned_ratings_hypotheses に保存し、Claude に集計を依頼した",
             "デスクトップPCでのLLM採点を実行し、結果をノートPCに戻した", "再投稿後にお礼と種明かしのメールを送った"])
    out = os.path.join(REV, "盲検評価_実施マニュアル（大槻先生用）.docx")
    doc.save(out); print("saved", out)


if __name__ == "__main__":
    rater_manual()
    author_guide()

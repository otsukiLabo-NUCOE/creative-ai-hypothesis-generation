# make_desktop_manual.py — デスクトップPC（GPU）でのLLM採点 手順書を生成する
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "desktop_judge_package", "手順書_デスクトップでのLLM採点.docx")
JP = "游ゴシック"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.0))
for name in ("Normal", "List Bullet", "List Number", "Heading 1", "Heading 2", "Title"):
    s = doc.styles[name]
    s.font.name = "Times New Roman"
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), JP)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Title", 18)):
    st = doc.styles[name]
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0, 0, 0)


def p(text="", bold=None, size=None, align=None, mono=False):
    para = doc.add_paragraph()
    if bold:
        para.add_run(bold).bold = True
    r = para.add_run(text)
    if size:
        r.font.size = Pt(size)
    if mono:
        r.font.name = "Consolas"
    if align:
        para.alignment = align
    return para


def bullet(text, bold=None, style="List Bullet"):
    para = doc.add_paragraph(style=style)
    if bold:
        para.add_run(bold).bold = True
    para.add_run(text)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def table(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(val); r.font.size = Pt(9.5)
            if i == 0:
                r.bold = True; shade(c, "D9D9D9")
    doc.add_paragraph()


def box(lines, fill):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.cell(0, 0); c.width = Cm(17); shade(c, fill); c.text = ""
    for k, (b, txt) in enumerate(lines):
        para = c.paragraphs[0] if k == 0 else c.add_paragraph()
        if b:
            r = para.add_run(b); r.bold = True; r.font.size = Pt(10)
        r = para.add_run(txt); r.font.size = Pt(10)
    doc.add_paragraph()


def code(text):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.cell(0, 0); c.width = Cm(17); shade(c, "F2F2F2"); c.text = ""
    r = c.paragraphs[0].add_run(text); r.font.name = "Consolas"; r.font.size = Pt(9.5)
    doc.add_paragraph()


doc.add_paragraph("デスクトップPCでのLLM採点 手順書", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
p("— GPU付きデスクトップで、原稿30本をLLMに採点させる手順（Claude・Python不要） —", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

box([("この手順書でやること", ""),
     ("", "・ノートPCで作った仮説30組と原稿30本（番号だけで、生成条件は分からない状態）を、デスクトップPCのGPUで大きめのLLMに採点させます。"),
     ("", "・採点は3種類です。①評価者と同じ仮説30組の順位づけ（30組×並び順3通り＝90回、短時間）、②原稿1本ずつの点数付け（30本×3回＝90回）、③原稿を2本ずつ比べる判定（90組×表示順2通り＝180回）。"),
     ("", "・できた結果フォルダをノートPCに戻せば、あとは Claude が集計して論文に反映します。"),
     ("所要時間の目安　", "準備（ダウンロードを含む）1〜2時間 ＋ 採点 2〜10時間（モデルとGPUによる。夜間に放置で構いません）")],
    "DDEBF7")

doc.add_heading("必要なもの", level=1)
for t in ["Windows 10 / 11 のデスクトップPC（NVIDIA の GPU 付き）",
          "インターネット接続（モデルのダウンロードに使います）",
          "Cドライブの空き容量 50GB 以上（1番の Llama-3.3-70B Q3_K_M の場合。2・3番なら 30GB 程度）",
          "USBメモリなど（このフォルダ desktop_judge_package を運ぶため。約1MB）"]:
    bullet(t)

# ------------------------------------------------------------------ STEP 1
doc.add_heading("STEP 1　GPUメモリを確認して、使うモデルを決める", level=1)
p("タスクマネージャー（Ctrl＋Shift＋Esc）→「パフォーマンス」→「GPU」を開き、「専用GPUメモリ」の値を確認します。"
  "「共有GPUメモリ」や、両者を合計した「GPUメモリ」の値ではありません（例：GPUメモリ32GB＝専用16GB＋共有16GBなら、16GBで判断します）。")
table([
    ["専用GPUメモリ", "使うモデル（番号）", "ファイルの大きさ", "採点時間の目安"],
    ["16GB〜24GB未満（今回のデスクトップは16GB）", "1：Llama-3.3-70B-Instruct Q3_K_M（第一候補）", "34.3GB", "6〜10時間"],
    ["12GB以上で、時間を短くしたい場合", "2：Qwen2.5-32B-Instruct Q4_K_M", "19.9GB", "2〜4時間"],
    ["8GB〜12GB未満", "3：Qwen2.5-14B-Instruct Q6_K", "12.1GB", "2〜3時間"],
    ["24GB 以上", "4：Llama-3.3-70B-Instruct Q4_K_M", "42.5GB", "6〜10時間"],
], [3.2, 6.8, 3.0, 4.0])
p("迷ったら、1つ小さいモデルを選んでください。途中でメモリ不足になるより確実です。", size=9.5)

# ------------------------------------------------------------------ STEP 2
doc.add_heading("STEP 2　フォルダをデスクトップPCにコピーする", level=1)
for t in [r"ノートPCの Revision1\desktop_judge_package フォルダを、USBメモリなどでデスクトップPCに運びます。",
          r"デスクトップPCの C:\ の直下にコピーし、フォルダ名を judge に変えます（C:\judge）。日本語や空白を含まない場所にしてください。",
          r"C:\judge の中に、manuscripts・prompts フォルダ、judge.ps1・run_judge.bat・download_model.bat・pairs.csv・hypothesis_sets.json があることを確認します。"]:
    bullet(t, style="List Number")

# ------------------------------------------------------------------ STEP 3
doc.add_heading("STEP 3　llama.cpp（GPU版）を用意する", level=1)
p("llama.cpp は、LLMを自分のPCで動かすための無料のソフトです。NVIDIA の GPU 用のものをダウンロードします。")
for t in ["ブラウザで https://github.com/ggml-org/llama.cpp/releases を開きます。",
          "一番上にある「b」で始まる番号のリリース（例：b10809）を開きます。「v0.5.0」のような v で始まるものはソースコードだけなので使いません。",
          "そのページの「Assets」から、次の2つのZIPをダウンロードします（12.4 の部分は、新しい番号でも構いません）。",
          ]:
    bullet(t, style="List Number")
code("llama-bXXXXX-bin-win-cuda-12.4-x64.zip      （本体）\ncudart-llama-bin-win-cuda-12.4-x64.zip       （CUDAランタイム）")
for t in [r"C:\judge の中に llama というフォルダを作り、2つのZIPの中身を両方ともそこに展開します。"
          r"C:\judge\llama\llama-server.exe があれば成功です（ZIPを展開したときに1段深いフォルダができた場合は、中身を llama 直下に移してください）。",
          "うまく動かない場合に備えて、「NVIDIA のドライバーが新しいこと」も確認してください（STEP 5 のトラブル対応を参照）。"]:
    bullet(t, style="List Number")

# ------------------------------------------------------------------ STEP 4
doc.add_heading("STEP 4　モデルをダウンロードする", level=1)
for t in [r"C:\judge\download_model.bat をダブルクリックします。",
          "黒い画面に STEP 1 で決めた番号（1〜4）を入力して Enter を押します。",
          "ダウンロードが始まります（30〜90分程度）。途中で止まった場合は、もう一度 download_model.bat を実行すれば続きから再開します。",
          r"「Finished: models\…」と表示されたら完了です。C:\judge\models にファイルができています。"]:
    bullet(t, style="List Number")
p("Windows が「WindowsによってPCが保護されました」と表示した場合は、「詳細情報」→「実行」を押してください。", size=9.5)

# ------------------------------------------------------------------ STEP 5
doc.add_heading("STEP 5　採点を実行する", level=1)
box([("先に必ず：", "「設定」→「システム」→「電源」で、スリープを「なし」にしてください。採点中にスリープすると止まります。")], "FCE4D6")
for t in [r"C:\judge\run_judge.bat をダブルクリックします。",
          "STEP 4 と同じ番号（1〜4）を入力して Enter を押します。",
          "モデルの読み込みに数分かかります。「llama-server (do not close)」という最小化されたウィンドウが開きますが、閉じないでください。",
          "「Server ready. Judging starts now.」と表示されたら採点が始まります。画面に進み具合が表示されます。",
          "最初に「[hypotheses 1/90] …」と仮説の順位づけが90回、続いて「[absolute 1/90] …」と原稿1本ずつの採点が90回、最後に「[pairwise 1/180] …」と2本ずつの比較が180回行われます。",
          "「Finished.」と表示されたら完了です。何かキーを押して画面を閉じます。"]:
    bullet(t, style="List Number")
p("途中で止まった・PCを再起動した場合：もう一度 run_judge.bat を実行し、同じ番号を選べば、終わった分は飛ばして続きから採点します。", size=9.5)

doc.add_heading("トラブル対応", level=2)
table([
    ["症状", "対応"],
    ["「The server did not start」と出る", r"C:\judge\results\（モデル名）\server.log を開きます。「cudart」「dll」などの語があれば、STEP 3 の cudart の ZIP を llama フォルダに展開し忘れています。「out of memory」「failed to allocate」があれば、メモリ不足です（次の行）"],
    ["メモリ不足（out of memory）", "1つ小さいモデル（番号を1つ大きく）を選び直してください。または run_judge.bat をメモ帳で開き、「set NGL=」を「set NGL=30」のように書き換えて GPU に載せる量を減らします（1番の Llama-3.3-70B Q3_K_M を専用16GBで使う場合は 30 程度から試す）"],
    ["極端に遅い（1回に10分以上）", "GPU が使われていない可能性があります。server.log に「CUDA」の語が無ければ、STEP 3 で CUDA 版ではないZIPを使っています。NVIDIA のドライバーを最新にしてください（nvidia-smi で表示される CUDA Version が 12.4 以上）"],
    ["「no valid answer, skipped」がたまに出る", "LLMの回答が形式どおりでなかったものです。数回なら問題ありません。最後に run_judge.bat をもう一度実行すると、その分だけ再挑戦します"],
    ["PowerShell の実行がブロックされる", "run_judge.bat から実行していれば通常は起きません。judge.ps1 を直接ダブルクリックしないでください"],
], [5.0, 12.0])

# ------------------------------------------------------------------ STEP 6
doc.add_heading("STEP 6　結果をノートPCに戻す", level=1)
for t in [r"C:\judge\results の中に、モデル名のフォルダ（例：Llama-3.3-70B-Instruct-Q4_K_M）ができています。",
          "その中に hypotheses（90ファイル）・absolute（90ファイル）・pairwise（180ファイル）のフォルダ、hypothesis_rankings.csv、absolute_ratings.csv、pairwise_results.csv、environment.txt、judge_log.txt があることを確認します。",
          r"このモデル名のフォルダごと、USBメモリでノートPCの Revision1\desktop_judge_results\ フォルダにコピーします（desktop_judge_results フォルダがなければ作ってください）。",
          "Claude に「デスクトップの採点結果を desktop_judge_results に入れたので集計して」と伝えます。"]:
    bullet(t, style="List Number")

doc.add_heading("注意", level=1)
for t in ["原稿は未発表の研究資料です。採点が終わったら、デスクトップPCの C:\\judge\\manuscripts は削除して構いません。",
          "結果フォルダには原稿の番号しか含まれず、どの条件で生成した原稿かは分かりません（対応表はノートPCにだけあります）。",
          "採点用のモデルや設定は、論文に記載するため変更しないでください（どうしても変える場合は Claude に伝えてください）。"]:
    bullet(t)

doc.add_heading("チェックリスト", level=1)
for t in ["専用GPUメモリを確認し、使うモデルの番号を決めた",
          r"C:\judge にフォルダをコピーした",
          r"C:\judge\llama\llama-server.exe がある（本体とcudartの2つのZIPを展開した）",
          "モデルのダウンロードが完了した（Finished と表示された）",
          "スリープを「なし」にしてから run_judge.bat を実行した",
          "Finished と表示された",
          r"結果フォルダをノートPCの Revision1\desktop_judge_results にコピーし、Claude に伝えた"]:
    para = doc.add_paragraph(); para.add_run("☐  ").font.size = Pt(12); para.add_run(t)

doc.save(OUT)
print("saved", OUT)

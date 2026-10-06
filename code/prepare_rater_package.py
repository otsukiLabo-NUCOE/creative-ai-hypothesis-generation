# prepare_rater_package.py
#
# Builds the material sent to the blinded expert raters (Appendix A):
#   ../rater_package/R<k>/manuscripts/M0xx.docx  (the 20 manuscripts assigned to rater k; blinded codes,
#                                                no generation metadata) + P00_practice.docx (practice)
#   ../rater_package/order_R<k>.csv              (rater-specific reading order)
#   ../results/controlled/_KEY_do_not_share.csv  (code -> condition, raters, shared; keep away from raters)
#
# Usage: python prepare_rater_package.py [--shared 10]

import argparse
import csv
import os
import random
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "results", "controlled", "manuscripts")
PRACTICE = os.path.join(HERE, "..", "results", "controlled", "practice", "practice.md")
PKG = os.path.join(HERE, "..", "rater_package")
KEY = os.path.join(HERE, "..", "results", "controlled", "_KEY_do_not_share.csv")
BLIND_SEED = 4591240


TEX_SYMBOLS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "varepsilon": "ε", "zeta": "ζ",
    "eta": "η", "theta": "θ", "vartheta": "θ", "iota": "ι", "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν",
    "xi": "ξ", "pi": "π", "rho": "ρ", "sigma": "σ", "tau": "τ", "phi": "φ", "varphi": "φ", "chi": "χ", "psi": "ψ",
    "omega": "ω", "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ",
    "Omega": "Ω", "Pi": "Π", "cdot": "·", "times": "×", "sum": "Σ", "prod": "Π", "in": "∈", "notin": "∉",
    "leq": "≤", "le": "≤", "geq": "≥", "ge": "≥", "neq": "≠", "approx": "≈", "sim": "∼", "propto": "∝",
    "infty": "∞", "partial": "∂", "nabla": "∇", "rightarrow": "→", "to": "→", "leftarrow": "←", "Rightarrow": "⇒",
    "mapsto": "↦", "mid": "|", "lVert": "‖", "rVert": "‖", "|": "‖", "forall": "∀", "exists": "∃", "cup": "∪",
    "cap": "∩", "subset": "⊂", "subseteq": "⊆", "oplus": "⊕", "otimes": "⊗", "circ": "∘", "ldots": "…",
    "cdots": "⋯", "dots": "…", "pm": "±", "int": "∫", "top": "T", "intercal": "T", "langle": "⟨", "rangle": "⟩",
    "log": "log", "exp": "exp", "max": "max", "min": "min", "arg": "arg", "argmax": "argmax", "argmin": "argmin",
    "sup": "sup", "inf": "inf", "lim": "lim", "sin": "sin", "cos": "cos", "tanh": "tanh", "det": "det",
    "quad": "  ", "qquad": "    ", ",": " ", ";": " ", ":": " ", "!": "", "{": "{", "}": "}", "_": "_", "%": "%",
}
DOUBLE_STRUCK = {"R": "ℝ", "N": "ℕ", "Z": "ℤ", "E": "𝔼", "P": "ℙ", "Q": "ℚ", "C": "ℂ"}


def _group(s, i):
    """Return (content, next_index) of a {...} group or single char starting at s[i]."""
    while i < len(s) and s[i] == " ":
        i += 1
    if i >= len(s):
        return "", i
    if s[i] == "{":
        depth, j = 0, i
        while j < len(s):
            depth += {"{": 1, "}": -1}.get(s[j], 0)
            if depth == 0:
                return s[i + 1:j], j + 1
            j += 1
        return s[i + 1:], len(s)
    if s[i] == "\\":
        m = re.match(r"\\([A-Za-z]+|.)", s[i:])
        return m.group(0), i + len(m.group(0))
    return s[i], i + 1


def tex_to_segments(tex):
    """LaTeX math -> list of (text, vert) with vert in {None, 'subscript', 'superscript'}."""
    s = tex.replace("\\left", "").replace("\\right", "").replace("\\displaystyle", "")
    # structural commands first
    for _ in range(6):
        s = re.sub(r"\\(?:d|t)?frac\s*", r"\\frac", s)
        m = re.search(r"\\frac", s)
        if not m:
            break
        a, j = _group(s, m.end()); b, k = _group(s, j)
        s = s[:m.start()] + f"({a})/({b})" + s[k:]
    s = re.sub(r"\\sqrt\s*\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\mathbb\s*\{?([A-Z])\}?", lambda mm: DOUBLE_STRUCK.get(mm.group(1), mm.group(1)), s)
    s = re.sub(r"\\(?:mathcal|mathbf|mathrm|mathit|boldsymbol|text|textbf|textit|operatorname|mathsf)\s*\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\(hat|bar|tilde|vec|dot)\s*\{?([A-Za-z])\}?",
               lambda mm: mm.group(2) + {"hat": "\u0302", "bar": "\u0304", "tilde": "\u0303", "vec": "\u20d7", "dot": "\u0307"}[mm.group(1)], s)
    s = re.sub(r"\\([A-Za-z]+|.)", lambda mm: TEX_SYMBOLS.get(mm.group(1), mm.group(1)), s)
    segs, buf, i = [], "", 0
    while i < len(s):
        ch = s[i]
        if ch in "_^":
            if buf:
                segs.append((buf, None)); buf = ""
            g, i = _group(s, i + 1)
            inner = "".join(t for t, _ in tex_to_segments(g)) if any(c in g for c in "_^\\") else g
            segs.append((inner.replace("{", "").replace("}", ""), "subscript" if ch == "_" else "superscript"))
            continue
        if ch not in "{}":
            buf += ch
        i += 1
    if buf:
        segs.append((buf, None))
    return segs


def _math(par, tex):
    for text, vert in tex_to_segments(tex):
        r = par.add_run(re.sub(r"\s{2,}", " ", text))
        r.font.name = "Cambria Math"
        if vert == "subscript":
            r.font.subscript = True
        elif vert == "superscript":
            r.font.superscript = True


def _inline(par, text):
    # $math$, **bold**, *italic*, `code`  ->  runs; strip other markdown markers
    for tok in re.split(r"(\$[^$]+\$|\\\([^)]*\\\)|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("$"):
            _math(par, tok[1:-1])
        elif tok.startswith("\\("):
            _math(par, tok[2:-2])
        elif tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def md_to_docx(md_text, out_path, header_code):
    doc = Document()
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.5))
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    hp = sec.header.paragraphs[0]
    hp.text = f"Manuscript {header_code}"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    in_math = False
    math_buf = []
    for line in md_text.splitlines():
        s = line.rstrip()
        one_line = re.fullmatch(r"\s*(\$\$|\\\[)(.+?)(\$\$|\\\])\s*", s)
        if one_line:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _math(p, one_line.group(2)); continue
        if s.strip() in ("$$", "\\[", "\\]"):
            if in_math:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _math(p, " ".join(math_buf))
                math_buf = []
            in_math = not in_math
            continue
        if in_math:
            math_buf.append(s.strip())
            continue
        if not s.strip() or re.fullmatch(r"-{3,}|\*{3,}", s.strip()):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).replace("**", ""), level=0 if level == 1 else min(level - 1, 3))
            continue
        m = re.match(r"^\s*[-*+]\s+(.*)", s)
        if m:
            _inline(doc.add_paragraph(style="List Bullet"), m.group(1)); continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", s)
        if m:
            _inline(doc.add_paragraph(style="List Number"), m.group(1)); continue
        if s.strip().startswith("|"):
            p = doc.add_paragraph(s.strip()); p.runs[0].font.name = "Consolas"; continue
        _inline(doc.add_paragraph(), s.strip())
    doc.save(out_path)


def main():
    """Partially overlapping design: 2 raters x 20 manuscripts. 10 manuscripts are rated by BOTH raters
    (inter-rater agreement), the remaining 20 are split 10/10, so every manuscript gets >= 1 human rating.
    Shared and unique sets are stratified by condition (shared 4/3/3; each rater ~7/7/6 per condition),
    so rater leniency is balanced across conditions. Each rater gets a folder R<k>/ with only their own
    manuscripts (+ the practice manuscript); the reading order is randomised per rater, and raters are
    not told which manuscripts are shared."""
    ap = argparse.ArgumentParser()
    ap.add_argument("--shared", type=int, default=10)
    args = ap.parse_args()
    files = sorted(f for f in os.listdir(SRC) if f.endswith(".md"))
    rng = random.Random(BLIND_SEED)
    rng.shuffle(files)
    code_of = {f: f"M{i:03d}" for i, f in enumerate(files, 1)}
    cond_of = {f: f.split("_")[0] for f in files}
    conds = sorted(set(cond_of.values()))
    arng = random.Random(BLIND_SEED + 7)
    # shared quota per condition (e.g. 4/3/3 for 10 shared over 3 conditions), random which gets the extra
    extra = arng.sample(conds, args.shared % len(conds))
    quota = {c: args.shared // len(conds) + (c in extra) for c in conds}
    assign, turn = {}, 0
    for c in conds:
        fs = [f for f in files if cond_of[f] == c]
        arng.shuffle(fs)
        for f in fs[:quota[c]]:
            assign[f] = (1, 2)
        for f in fs[quota[c]:]:                 # deal the rest alternately, continuing across conditions
            assign[f] = (turn % 2 + 1,)
            turn += 1
    with open(KEY, "w", newline="", encoding="utf-8") as k:
        kw = csv.writer(k)
        kw.writerow(["code", "original_file", "condition", "raters", "shared"])
        for f in files:
            kw.writerow([code_of[f], f, cond_of[f], ";".join(f"R{r}" for r in assign[f]), int(len(assign[f]) == 2)])
    for r in (1, 2):
        out = os.path.join(PKG, f"R{r}", "manuscripts")
        os.makedirs(out, exist_ok=True)
        mine = [f for f in files if r in assign[f]]
        for f in mine:
            md_to_docx(open(os.path.join(SRC, f), encoding="utf-8").read(),
                       os.path.join(out, code_of[f] + ".docx"), code_of[f])
        if os.path.exists(PRACTICE):
            md_to_docx(open(PRACTICE, encoding="utf-8").read(), os.path.join(out, "P00_practice.docx"), "P00 (practice)")
        order = sorted(code_of[f] for f in mine)
        random.Random(BLIND_SEED + r).shuffle(order)          # rater-specific reading order
        with open(os.path.join(PKG, f"order_R{r}.csv"), "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh); w.writerow(["position", "code"])
            for pos, c in enumerate(order, 1):
                w.writerow([pos, c])
        print(f"R{r}: {len(mine)} manuscripts "
              f"({', '.join(f'{c}={sum(cond_of[f] == c for f in mine)}' for c in sorted(set(cond_of.values())))})")
    print(f"key -> {KEY}")


if __name__ == "__main__":
    main()
